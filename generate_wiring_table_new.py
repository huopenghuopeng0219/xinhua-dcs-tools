import openpyxl
import os
import sys
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, GradientFill, Color
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter


class XlsWrapper:
    """xlrd工作表的包装器，提供与openpyxl兼容的接口"""
    def __init__(self, xls_sheet):
        self.xls_sheet = xls_sheet
        self.max_row = xls_sheet.nrows
        self.max_column = xls_sheet.ncols
        self.merged_cells = []
        
        # 提取合并单元格信息
        for merge_range in xls_sheet.merged_cells:
            class MergeObj:
                def __init__(self, r1, r2, c1, c2):
                    self.min_row = r1
                    self.max_row = r2
                    self.min_col = c1
                    self.max_col = c2
            self.merged_cells.append(MergeObj(merge_range[0], merge_range[1], merge_range[2], merge_range[3]))
    
    def cell(self, row, column):
        """获取单元格对象（1-based索引）"""
        # xlrd是0-based，转换为1-based
        value = self.xls_sheet.cell_value(row - 1, column - 1)
        # 返回一个带value属性的对象
        class CellObj:
            def __init__(self, v):
                self.value = v
        return CellObj(value)


class WiringTableGenerator:
    """接线总表生成器 - 从卡件排布工作表生成接线总表"""

    # 卡件排布中的目标列：(列号, 列名, P值来源列)
    # P值来源列：I=9, AN=40, 可根据实际文件调整
    TARGET_COLUMNS = [(12, "L", 9), (19, "S", 9), (26, "Z", 9),
                      (36, "AJ", 40), (43, "AQ", 40), (50, "AX", 40)]

    # 信号类型与通道数映射
    SIGNAL_TYPE_16 = {"DI", "DO", "SOE"}  # 16通道
    SIGNAL_TYPE_8 = {"AI", "AO", "PI", "RTD", "TC"}  # 8通道

    # 测点类型映射
    POINT_TYPE_MAP = {
        "DI": "DI", "SOE": "SOE", "PI": "PI", "DO": "DO",
        "AI": "AI", "AO": "AO", "RTD": "Pt100", "TC": "TC-K"
    }

    # 表头定义（A-W列，共23列）
    HEADERS = [
        "测点名", "中文描述", "DPU", "站号", "卡号", "卡通道号", "卡件类型",
        "测点类型", "特征字", "单位", "下限", "上限", "端子柜号", "正反面",
        "端子板号", "接线板号", "AI是否供电\nDO是否常闭", "端子(+)",
        "端子(-)", "端子(-)", "号码管1", "号码管2", "号码管3"
    ]
    COL_WIDTHS = [20, 30, 5, 5, 5, 5, 8, 8, 8, 5, 5, 5, 15, 5, 5, 5, 12, 10, 10, 10, 25, 25, 25]

    def __init__(self, file_path):
        self.file_path = file_path
        # 根据文件扩展名选择读取方式
        if file_path.lower().endswith('.xls'):
            # .xls文件用xlrd读取
            import xlrd
            xls_wb = xlrd.open_workbook(file_path, formatting_info=True)
            self.xls_workbook = xls_wb
            self.workbook = None
            # 包装所有工作表
            self.wrapped_sheets = {}
            for name in xls_wb.sheet_names():
                self.wrapped_sheets[name] = XlsWrapper(xls_wb.sheet_by_name(name))
        else:
            # .xlsx/.xlsm文件用openpyxl读取
            self.workbook = openpyxl.load_workbook(file_path, data_only=False)
            self.xls_workbook = None
            self.wrapped_sheets = {}

    def generate(self, output_path):
        """生成接线总表"""
        # 收集所有数据行
        all_rows = []
        
        # 获取工作表列表
        if self.xls_workbook:
            sheet_names = self.xls_workbook.sheet_names()
        else:
            sheet_names = self.workbook.sheetnames

        for sheet_name in sheet_names:
            sheet_lower = sheet_name.lower()
            if not (sheet_lower.endswith("-az") or sheet_lower.startswith("卡件排布")):
                continue

            if self.xls_workbook:
                ws = self.wrapped_sheets[sheet_name]
            else:
                ws = self.workbook[sheet_name]
            
            rows = self._parse_sheet(ws, sheet_name)
            all_rows.extend(rows)

        # 过滤不需要的行
        filtered_rows = self._filter_rows(all_rows)

        # 写入Excel
        self._write_output(output_path, filtered_rows)

        return len(filtered_rows)

    def _parse_sheet(self, ws, sheet_name):
        """解析单个工作表"""
        rows = []

        for col_num, col_name, p_col in self.TARGET_COLUMNS:
            # 获取DPU-站号
            cell_value = ws.cell(row=20, column=col_num).value
            if not cell_value:
                continue

            dpu, station = self._parse_dpu_station(cell_value)

            # 跳过模板（DPU=00）
            if dpu == "00":
                continue

            # 获取端子柜号
            cabinet = ws.cell(row=72, column=34).value or ""

            # 获取卡件前缀
            card_prefix = ws.cell(row=18, column=col_num).value or ""

            # 扫描数据块
            current_row = 22
            block_index = 0

            while current_row <= 63:
                block_data = self._parse_block(ws, current_row, col_num)
                if block_data:
                    card_no, signal_type, channel_count, card_num = block_data

                    # 跳过空壳
                    if "空壳" in card_no:
                        current_row += 4
                        continue

                    # 确定实际通道数（DI/DO/SOE=16，其他=8）
                    if signal_type in self.SIGNAL_TYPE_16:
                        actual_channels = 16
                    else:
                        actual_channels = 8

                    # 获取P值来源
                    if signal_type == "DO":
                        # DO信号从卡件排布读取
                        p_row = 23 + block_index * 4
                        p_value = ws.cell(row=p_row, column=p_col).value or ""
                    else:
                        # 非DO信号P=O
                        p_value = f"{card_prefix}{card_num}"

                    # 生成每一行
                    o_value = f"{card_prefix}{card_num}"

                    for ch in range(actual_channels):
                        row_data = self._build_row(
                            dpu=dpu,
                            station=station,
                            block_no=card_num,
                            channel=ch,
                            signal_type=signal_type,
                            card_prefix=card_prefix,
                            o_value=o_value,
                            p_value=p_value,
                            cabinet=cabinet
                        )
                        rows.append(row_data)

                    block_index += 1
                    current_row += 4
                else:
                    current_row += 1

        return rows

    def _parse_block(self, ws, row, col):
        """解析数据块（支持合并单元格和单个单元格）"""
        # 先尝试获取合并单元格
        merge = self._get_merge_area(ws, row, col)
        if merge:
            min_col, max_col, min_row, max_row = merge
            merge_cols = max_col - min_col + 1
            merge_rows = max_row - min_row + 1
            
            if merge_cols == 3 and merge_rows == 4 and min_col == col:
                cell_value = ws.cell(row=min_row, column=min_col).value
                if cell_value:
                    return self._parse_block_content(cell_value)
        
        # 尝试单个单元格（xls格式）
        cell_value = ws.cell(row=row, column=col).value
        if cell_value and "\n" in str(cell_value):
            return self._parse_block_content(cell_value)
        
        return None
    
    def _parse_block_content(self, cell_value):
        """解析单元格内容"""
        lines = str(cell_value).strip().split("\n")
        if len(lines) < 3:
            return None
        
        card_no = lines[0].strip()
        signal_type = lines[1].strip().upper()
        
        try:
            channel_count = int(lines[2].strip())
        except ValueError:
            return None
        
        # 验证信号类型
        if signal_type not in self.SIGNAL_TYPE_16 and signal_type not in self.SIGNAL_TYPE_8:
            return None
        
        # 提取卡号数字
        import re
        num_match = re.search(r'(\d+)', card_no)
        card_num = int(num_match.group(1)) if num_match else 0
        
        return card_no, signal_type, channel_count, card_num

    def _build_row(self, dpu, station, block_no, channel, signal_type,
                   card_prefix, o_value, p_value, cabinet):
        """构建一行数据"""
        # 测点类型
        point_type = self.POINT_TYPE_MAP.get(signal_type, "")

        # 特征字
        feature = "DCS"

        # 单位和上下限
        unit, lower, upper = self._get_unit_limits(signal_type)

        # 正反面
        front_back = self._get_front_back(card_prefix)

        # 端子号计算
        r_value = self._calc_terminal_1(signal_type, channel, None)
        s_value = self._calc_terminal_2(signal_type, channel, None)
        t_value = self._calc_terminal_3(signal_type, channel)

        return {
            "A": "",  # 测点名（用户填写）
            "B": "",  # 中文描述（用户填写）
            "C": dpu,
            "D": station,
            "E": block_no,
            "F": channel,
            "G": signal_type,
            "H": point_type,
            "I": feature,
            "J": unit,
            "K": lower,
            "L": upper,
            "M": cabinet,
            "N": front_back,
            "O": o_value,
            "P": p_value,
            "Q": "",  # AI是否供电/DO是否常闭
            "R": r_value,
            "S": s_value,
            "T": t_value,
            "U": "",  # 号码管1
            "V": "",  # 号码管2
            "W": "",  # 号码管3
        }

    def _get_unit_limits(self, signal_type):
        """获取单位和上下限"""
        if signal_type == "TC":
            return "℃", 0, 1232
        elif signal_type == "RTD":
            return "℃", -50, 500
        return "", "", ""

    def _get_front_back(self, card_prefix):
        """根据卡件前缀判断正反面"""
        prefix = str(card_prefix).upper()
        if any(c in prefix for c in ["A", "B", "C"]):
            return "正"
        elif any(c in prefix for c in ["D", "E", "F"]):
            return "反"
        return ""

    def _calc_terminal_1(self, signal_type, channel, q_value):
        """计算新端子1+（R列）"""
        q_is_1 = (q_value == 1)

        if signal_type in ("DI", "SOE"):
            return channel + 1
        elif signal_type in ("PI", "AO"):
            return 2 * channel + 1
        elif signal_type == "DO":
            if channel < 8:
                return 2 * channel + 2 if q_is_1 else 2 * channel + 1
            else:
                return 2 * channel + 18 if q_is_1 else 2 * channel + 17
        elif signal_type == "AI":
            return 2 * channel + 20 if q_is_1 else 2 * channel + 1
        elif signal_type in ("RTD", "TC"):
            return 2 * channel + 1
        return ""

    def _calc_terminal_2(self, signal_type, channel, q_value):
        """计算新端子2-（S列）"""
        q_is_1 = (q_value == 1)

        if signal_type in ("DI", "SOE"):
            return channel + 19
        elif signal_type in ("PI", "AO"):
            return 2 * channel + 19
        elif signal_type == "DO":
            if channel < 8:
                return 2 * channel + 18 if q_is_1 else 2 * channel + 17
            else:
                return 2 * channel + 34 if q_is_1 else 2 * channel + 33
        elif signal_type == "AI":
            return 2 * channel + 1 if q_is_1 else 2 * channel + 19
        elif signal_type in ("RTD", "TC"):
            return 2 * channel + 2
        return ""

    def _calc_terminal_3(self, signal_type, channel):
        """计算新端子3-（T列）"""
        if signal_type == "RTD":
            return 2 * channel + 19
        return ""

    def _get_merge_area(self, ws, row, col):
        """获取合并单元格区域"""
        for merged in ws.merged_cells:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                return (merged.min_col, merged.max_col, merged.min_row, merged.max_row)
        return None

    def _parse_dpu_station(self, value):
        """解析DPU-站号"""
        s = str(value)
        if "-" in s:
            parts = s.split("-", 1)
            return parts[0], parts[1]
        return s, ""

    def _filter_rows(self, rows):
        """过滤不需要的行（参照VBA逻辑）"""
        filtered = []
        for row in rows:
            g = row["G"]
            f = row["F"]

            # G为空则删除
            if not g:
                continue

            # AI/PI/AO/RTD/TC 且 F>=8 则删除
            if g in ("AI", "PI", "AO", "RTD", "TC") and f >= 8:
                continue

            filtered.append(row)
        return filtered

    def _write_output(self, output_path, rows):
        """写入输出文件"""
        if output_path.lower().endswith('.xls') and not output_path.lower().endswith('.xlsx'):
            self._write_xls(output_path, rows)
        else:
            self._write_xlsx(output_path, rows)

    def _write_xlsx(self, output_path, rows):
        """写入xlsx文件"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "接线总表"

        # 写表头
        for col, header in enumerate(self.HEADERS, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = self.COL_WIDTHS[col - 1]
        ws.row_dimensions[1].height = 32

        # 写数据
        keys = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L",
                "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W"]

        for row_idx, row_data in enumerate(rows, 2):
            # 写入静态值 (C-Q)
            for col_idx in range(3, 18):  # C=3 到 Q=17
                key = keys[col_idx - 1]
                ws.cell(row=row_idx, column=col_idx, value=row_data[key])

            # 写入R/S/T/U/V/W公式
            r = row_idx
            # R列: 端子1+
            formula_r = f'=IF(OR(G{r}="DI",G{r}="SOE"),F{r}+1,IF(OR(G{r}="PI",G{r}="AO"),2*F{r}+1,IF(G{r}="DO",IF(F{r}<8,IF(Q{r}=1,2*F{r}+2,2*F{r}+1),IF(Q{r}=1,2*F{r}+18,2*F{r}+17)),IF(G{r}="AI",IF(Q{r}=1,2*F{r}+20,2*F{r}+1),IF(OR(G{r}="RTD",G{r}="TC"),2*F{r}+1,"")))))'
            ws.cell(row=r, column=18, value=formula_r)

            # S列: 端子2-
            formula_s = f'=IF(OR(G{r}="DI",G{r}="SOE"),F{r}+19,IF(OR(G{r}="PI",G{r}="AO"),2*F{r}+19,IF(G{r}="DO",IF(F{r}<8,IF(Q{r}=1,2*F{r}+18,2*F{r}+17),IF(Q{r}=1,2*F{r}+34,2*F{r}+33)),IF(G{r}="AI",IF(Q{r}=1,2*F{r}+1,2*F{r}+19),IF(OR(G{r}="RTD",G{r}="TC"),2*F{r}+2,"")))))'
            ws.cell(row=r, column=19, value=formula_s)

            # T列: 端子3-
            formula_t = f'=IF(G{r}="RTD",2*F{r}+19,"")'
            ws.cell(row=r, column=20, value=formula_t)

            # U列: 号码管1
            formula_u = f'=IF(ISBLANK(A{r}),"",A{r}&"/"&P{r}&"/"&R{r})'
            ws.cell(row=r, column=21, value=formula_u)

            # V列: 号码管2
            formula_v = f'=IF(ISBLANK(A{r}),"",A{r}&"/"&P{r}&"/"&S{r})'
            ws.cell(row=r, column=22, value=formula_v)

            # W列: 号码管3
            formula_w = f'=IF(G{r}="RTD",IF(ISBLANK(A{r}),"",A{r}&"/"&P{r}&"/"&T{r}),"")'
            ws.cell(row=r, column=23, value=formula_w)

        # 设置格式
        self._apply_format_openpyxl(wb, ws, len(rows))
        
        # 冻结窗格（冻结第1行）
        ws.freeze_panes = "A2"
        
        # 打印标题（每页打印第1行）
        ws.print_title_rows = "1:1"
        
        # 页面设置（用try避免兼容性问题）
        try:
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_margins.left = 0.5
            ws.page_margins.right = 0.5
            ws.page_margins.top = 0.5
            ws.page_margins.bottom = 1
            ws.page_margins.header = 0.5
            ws.page_margins.footer = 0.5
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.oddFooter.center.text = "第 &P 页，共 &N 页"
        except Exception as e:
            print(f"警告：部分页面设置跳过 - {e}")
        
        # 保存文件
        wb.save(output_path)

    def _write_xls(self, output_path, rows):
        """写入xls文件（WPS兼容）"""
        import xlwt
        
        wb = xlwt.Workbook(encoding='utf-8')
        wb.codepage = 0x04B0  # UTF-16
        ws = wb.add_sheet('接线总表')
        
        # 表头样式
        header_style = xlwt.easyxf(
            'font: bold on, height 220;'
            'alignment: horiz centre, vert centre, wrap on;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 数据样式
        data_style = xlwt.easyxf(
            'font: height 220;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 左对齐样式
        left_style = xlwt.easyxf(
            'font: height 220;'
            'alignment: horiz left, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 黄色高亮样式
        yellow_style = xlwt.easyxf(
            'font: height 220;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
            'pattern: pattern solid, fore_colour yellow;'
        )
        
        yellow_left_style = xlwt.easyxf(
            'font: height 220;'
            'alignment: horiz left, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
            'pattern: pattern solid, fore_colour yellow;'
        )
        
        # 写表头
        for col, header in enumerate(self.HEADERS):
            ws.write(0, col, header, header_style)
            ws.col(col).width = self.COL_WIDTHS[col] * 256
        
        # 写数据
        keys = ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L",
                "M", "N", "O", "P", "Q", "R", "S", "T"]
        
        for row_idx, row_data in enumerate(rows, 1):
            # 检查是否需要高亮
            f_val = row_data.get("F")
            g_val = row_data.get("G")
            need_highlight = False
            if f_val is not None and g_val:
                cond1 = (f_val == 15) and (g_val in ("DI", "DO", "SOE"))
                cond2 = (f_val == 7) and (g_val in ("AI", "AO", "TC", "RTD", "PI"))
                need_highlight = cond1 or cond2
            
            for col_idx, key in enumerate(keys):
                value = row_data[key]
                if value == "":
                    value = None
                
                # 选择样式
                if need_highlight:
                    style = yellow_left_style if col_idx < 2 else yellow_style
                else:
                    style = left_style if col_idx < 2 else data_style
                
                ws.write(row_idx, col_idx, value, style)
        
        wb.save(output_path)
        print(f"已生成xls文件: {output_path}")

    def _fix_xml_for_wps(self, file_path):
        """修复XML以兼容WPS"""
        import zipfile
        import re
        
        temp_file = file_path + ".tmp"
        
        with zipfile.ZipFile(file_path, 'r') as zin:
            files = {}
            for name in zin.namelist():
                files[name] = zin.read(name)
            
            # 修复styles.xml
            if 'xl/styles.xml' in files:
                styles = files['xl/styles.xml'].decode('utf-8')
                # 替换dxf为正确的格式
                styles = re.sub(
                    r'<dxfs.*?</dxfs>',
                    '<dxfs count="1"><dxf><fill><patternFill patternType="solid"><fgColor rgb="00FFFF00" /><bgColor rgb="00FFFF00" /></patternFill></fill></dxf></dxfs>',
                    styles, flags=re.DOTALL
                )
                files['xl/styles.xml'] = styles.encode('utf-8')
            
            # 修复sheet1.xml的条件格式
            if 'xl/worksheets/sheet1.xml' in files:
                sheet = files['xl/worksheets/sheet1.xml'].decode('utf-8')
                # 移除openpyxl生成的条件格式
                sheet = re.sub(r'<conditionalFormatting.*?</conditionalFormatting>', '', sheet, flags=re.DOTALL)
                # 添加正确的条件格式
                cf_xml = '''<conditionalFormatting sqref="A2:T3153">
<cfRule type="expression" priority="1" dxfId="0">
<formula>=OR(AND($F1=15,OR($G1="DI",$G1="DO",$G1="SOE")),AND($F1=7,OR($G1="AI",$G1="AO",$G1="TC",$G1="RTD",$G1="PI")))</formula>
</cfRule>
</conditionalFormatting>'''
                sheet = sheet.replace('</worksheet>', cf_xml + '</worksheet>')
                files['xl/worksheets/sheet1.xml'] = sheet.encode('utf-8')
            
            # 写入临时文件
            with zipfile.ZipFile(temp_file, 'w', zipfile.ZIP_DEFLATED) as zout:
                for name, data in files.items():
                    zout.writestr(name, data)
        
        # 替换原文件
        import os
        try:
            os.replace(temp_file, file_path)
        except PermissionError:
            # 文件被占用，尝试删除后重试
            try:
                os.remove(file_path)
                os.replace(temp_file, file_path)
            except:
                # 如果还是失败，用新文件名
                new_path = file_path.replace('.xlsx', '_new.xlsx')
                os.replace(temp_file, new_path)
                print(f"警告：原文件被占用，已保存为 {os.path.basename(new_path)}")

    def _apply_format_openpyxl(self, wb, ws, data_rows):
        """应用表格格式"""
        if data_rows < 1:
            return

        last_row = data_rows + 1
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )

        # 字体和对齐
        for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=23):
            for cell in row:
                cell.font = Font(size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=True)

        # 表头居中
        for col in range(1, 24):
            ws.cell(row=1, column=col).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

        # A-B列左对齐
        for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        # C-W列居中对齐
        for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=3, max_col=23):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # 添加边框
        for row in ws.iter_rows(min_row=1, max_row=last_row, min_col=1, max_col=23):
            for cell in row:
                cell.border = thin_border

        # 条件格式：直接给符合条件的行涂色（WPS兼容）
        print("正在设置条件格式...")
        for row in range(2, last_row + 1):
            f_val = ws.cell(row=row, column=6).value
            g_val = ws.cell(row=row, column=7).value

            if f_val is not None and g_val:
                # DI/DO/SOE 且 F=15，或 AI/AO/TC/RTD/PI 且 F=7
                cond1 = (f_val == 15) and (g_val in ("DI", "DO", "SOE"))
                cond2 = (f_val == 7) and (g_val in ("AI", "AO", "TC", "RTD", "PI"))

                if cond1 or cond2:
                    for col in range(1, 24):
                        ws.cell(row=row, column=col).fill = PatternFill(
                            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
                        )


class ConfigGenerator:
    """转组态文件生成器 - 从接线总表生成转组态xls文件"""
    
    # 接线总表列 -> 转组态IO表列的映射（A-P，共16列）
    COLUMN_MAPPING = {
        0: 0,   # A -> A (测点名)
        1: 1,   # B -> B (中文描述)
        2: 2,   # C -> C (DPU)
        3: 3,   # D -> D (站号)
        4: 4,   # E -> E (卡号)
        5: 5,   # F -> F (卡通道号)
        6: 6,   # G -> G (卡件类型)
        7: 7,   # H -> H (测点类型)
        8: 8,   # I -> I (特征字)
        9: 9,   # J -> J (单位)
        10: 10, # K -> K (下限)
        11: 11, # L -> L (上限)
        12: 12, # M -> M (端子柜号)
        13: 13, # N -> N (正反面)
        14: 14, # O -> O (端子板号)
        15: 15, # P -> P (DO端子板号)
    }
    
    # IO表表头（A-Z，26列）
    IO_HEADERS = [
        "测点名", "中文描述", "DPU", "站号", "卡号", "卡通道号", "卡件类型",
        "测点类型", "特征字", "单位", "下限", "上限", "端子柜号", "正反面",
        "端子板号", "DO端子板号", "开关报警", "机柜", "测点组", "供电",
        "DO类型", "端子1", "端子2", "端子3", "KKS", "系统"
    ]
    
    # DEVICE表表头（A-V，22列）
    DEVICE_HEADERS = [
        "VC", "描述", "DPU", "开指令", "关指令", "开到位", "关到位",
        "远方控制", "电源消失", "故障", "其他1", "其他2", "其他3", "其他4",
        "其他5", "其他6", "页号", "位置", "系统", "子图图", "操作面板", "KKS"
    ]
    
    # ESMA表表头（A-M，13列）
    ESMA_HEADERS = [
        "设定", "描述", "DPU", "指令", "反馈", "在手动", "切手动",
        "页号", "位置", "系统", "子图图", "操作面板", "KKS"
    ]
    
    def __init__(self, wiring_table_path):
        self.wiring_table_path = wiring_table_path
    
    def generate(self, output_path):
        """生成转组态文件"""
        import xlwt
        
        # 读取接线总表
        wb_in = openpyxl.load_workbook(self.wiring_table_path, data_only=True)
        ws_in = wb_in['接线总表']
        
        # 创建输出文件（用xlwt生成xls，ODBC驱动兼容）
        wb_out = xlwt.Workbook(encoding='utf-8')
        wb_out.codepage = 0x04B0  # UTF-16
        
        # 创建IO表
        self._create_io_sheet(wb_out, ws_in)
        
        # 创建DEVICE表（只有表头）
        self._create_device_sheet(wb_out)
        
        # 创建ESMA表（只有表头）
        self._create_esma_sheet(wb_out)
        
        # 保存
        wb_out.save(output_path)
        print(f"转组态文件已生成: {output_path}")
    
    def _create_io_sheet(self, wb_out, ws_in):
        """创建IO表"""
        import xlwt
        
        ws_out = wb_out.add_sheet('IO')
        
        # 定义样式
        # 表头样式：宋体12号，居中，四边框线
        header_style = xlwt.easyxf(
            'font: name 宋体, height 240;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 数据样式：宋体10号，居中，四边框线
        data_style = xlwt.easyxf(
            'font: name 宋体, height 200;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # A/B/J列（文本格式）
        text_style = xlwt.easyxf(
            'font: name 宋体, height 200;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;',
            num_format_str='@'
        )
        
        # K/L列（上下限）样式：数值格式，2位小数
        number_style = xlwt.easyxf(
            'font: name 宋体, height 200;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;',
            num_format_str='0.00'
        )
        
        # 写表头（A-Z，26列）
        for col, header in enumerate(self.IO_HEADERS):
            ws_out.write(0, col, header, header_style)
            # 设置列宽
            if col < 2:
                ws_out.col(col).width = 256 * 18  # 测点名、中文描述
            else:
                ws_out.col(col).width = 256 * 10
        
        # 写数据行
        target_row = 1
        for row in range(2, ws_in.max_row + 1):
            # 检查是否有数据
            if ws_in.cell(row=row, column=1).value is None and ws_in.cell(row=row, column=3).value is None:
                continue
            
            # 映射列（A-P，共16列有数据）
            for src_col, dst_col in self.COLUMN_MAPPING.items():
                value = ws_in.cell(row=row, column=src_col + 1).value
                if value is None:
                    value = ""
                
                # 特殊处理：DPU和站号转为数字
                if dst_col in [2, 3]:  # C, D列
                    try:
                        value = float(value) if value != "" else ""
                    except:
                        pass
                
                # 特殊处理：卡号和卡通道号转为数字
                if dst_col in [4, 5]:  # E, F列
                    try:
                        value = float(value) if value != "" else ""
                    except:
                        pass
                
                # 特殊处理：下限和上限转为数字
                if dst_col in [10, 11]:  # K, L列
                    try:
                        value = float(value) if value != "" else ""
                    except:
                        pass
                
                # 选择样式
                if dst_col in [0, 1, 9]:  # A/B/J列（文本格式）
                    style = text_style
                elif dst_col in [10, 11]:  # K/L列（上下限，数值格式）
                    style = number_style
                else:
                    style = data_style
                
                ws_out.write(target_row, dst_col, value, style)
            
            # Q-Z列留空（需要用户填写或后续处理）
            for col in range(16, 26):
                ws_out.write(target_row, col, "", data_style)
            
            target_row += 1
        
        print(f"IO表: 共 {target_row - 1} 行数据")
    
    def _create_device_sheet(self, wb_out):
        """创建DEVICE表（只有表头）"""
        import xlwt
        
        ws_out = wb_out.add_sheet('DEVICE')
        
        # 表头样式
        header_style = xlwt.easyxf(
            'font: name 宋体, height 240;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 写表头（A-V，22列）
        for col, header in enumerate(self.DEVICE_HEADERS):
            ws_out.write(0, col, header, header_style)
            ws_out.col(col).width = 256 * 12
        
        print("DEVICE表: 已创建（仅表头）")
    
    def _create_esma_sheet(self, wb_out):
        """创建ESMA表（只有表头）"""
        import xlwt
        
        ws_out = wb_out.add_sheet('ESMA')
        
        # 表头样式
        header_style = xlwt.easyxf(
            'font: name 宋体, height 240;'
            'alignment: horiz centre, vert centre;'
            'borders: left thin, right thin, top thin, bottom thin;'
        )
        
        # 写表头（A-M，13列）
        for col, header in enumerate(self.ESMA_HEADERS):
            ws_out.write(0, col, header, header_style)
            ws_out.col(col).width = 256 * 12
        
        print("ESMA表: 已创建（仅表头）")


def main():
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    else:
        file_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "1234.xlsm")

    if not os.path.exists(file_path):
        print(f"错误: 文件不存在 - {file_path}")
        print("用法: python generate_wiring_table_new.py <Excel文件路径> [xls|xlsx|转组态]")
        return

    # 检查参数
    output_format = "xlsx"
    generate_config = False
    
    if len(sys.argv) > 2:
        arg = sys.argv[2].lower()
        if arg == "xls":
            output_format = "xls"
        elif arg == "转组态":
            generate_config = True

    print(f"正在处理: {file_path}")

    if generate_config:
        # 生成转组态文件
        output_path = os.path.join(os.path.dirname(file_path), "转组态.xls")
        generator = ConfigGenerator(file_path)
        generator.generate(output_path)
    else:
        # 生成接线总表
        print(f"输出格式: {output_format}")
        generator = WiringTableGenerator(file_path)
        ext = f".{output_format}"
        output_path = os.path.join(os.path.dirname(file_path), f"接线总表{ext}")
        count = generator.generate(output_path)
        print(f"生成完成！共 {count} 行数据")
        print(f"保存至: {output_path}")


if __name__ == "__main__":
    main()
