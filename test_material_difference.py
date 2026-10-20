from collections import defaultdict
from pathlib import Path
from tempfile import TemporaryDirectory

import openpyxl
from openpyxl.styles import Font

from material_difference import generate_difference


def test_sample():
    source = next(Path(__file__).parent.glob('*机柜设计*.xlsx'))
    target = source.with_name(source.stem.replace('机柜设计', '差异表') + '-自动测试.xlsx')
    _, count = generate_difference(source, target)
    original = openpyxl.load_workbook(source, data_only=True)
    actual = defaultdict(int)
    source_rows = 0
    mx_sheets = 0
    for ws in original.worksheets:
        if not ws.title.endswith('-MX'):
            continue
        mx_sheets += 1
        marker = next(r for r in range(1, ws.max_row + 1)
                      if str(ws[f'S{r}'].value or '').replace(' ', '') == '整件')
        for row in range(marker + 1, ws.max_row + 1):
            code = ws[f'J{row}'].value
            if code:
                source_rows += 1
                actual[str(code).strip()] += ws[f'AM{row}'].value

    output = openpyxl.load_workbook(target)
    ws = output.active
    assert output.sheetnames == ['机柜差异表']
    assert count == len(actual) == 54
    assert mx_sheets == 10
    assert source_rows == 133
    assert sum(actual.values()) == 843
    assert ws.max_column == 52
    assert ws.max_row == 111
    assert ws['A1'].value == 'XK03N2026703青岛平度兴平热电1x5万千瓦背压机组热电联产项目-差异表'
    assert 'A1:AZ1' in {str(m) for m in ws.merged_cells.ranges}
    assert ws['AL2'].value == '合同数量'
    assert ws['AU2'].value == '数量差异'
    assert ws['AZ2'].value == '原因'
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert 'AZ2:AZ3' in merged
    assert ws.sheet_view.view == 'normal'
    assert ws.sheet_properties.pageSetUpPr.autoPageBreaks is False
    assert ws.print_area == ''
    assert all(cell.value not in {'整    件', '整件', '设计编制'}
               for row in ws.iter_rows() for cell in row)
    for index, (code, quantity) in enumerate(actual.items()):
        row = 4 + 2 * index
        assert ws[f'A{row}'].value == index + 1
        assert ws[f'E{row}'].value == code
        assert ws[f'AH{row}'].value == quantity
        assert ws[f'AL{row}'].value is None
        assert ws[f'AU{row}'].value == f'=IF(AL{row}="","",AH{row}-AL{row})'
        assert ws[f'AZ{row}'].value is None
        assert f'AZ{row}:AZ{row + 1}' in merged

    allowed = {None, 'thin'}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=52):
        for cell in row:
            assert {cell.border.left.style, cell.border.right.style,
                    cell.border.top.style, cell.border.bottom.style} <= allowed
    for address in ('A1', 'A2', 'E4', 'AZ2', 'AZ4'):
        cell = ws[address]
        assert {cell.border.left.style, cell.border.right.style,
                cell.border.top.style, cell.border.bottom.style} == {'thin'}



def test_xlsx_keeps_original_layout_path():
    with TemporaryDirectory(dir=Path(__file__).parent) as directory:
        source = Path(directory) / "测试项目-机柜设计-V0.xlsx"
        target = Path(directory) / "测试项目-差异表-V0.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "01-MX"
        sheet.sheet_format.defaultColWidth = 1.6
        sheet.column_dimensions["F"].width = 40
        for column, name in {8: "代号", 9: "名称", 39: "数量", 43: "备注", 47: "更改"}.items():
            sheet.cell(2, column, name)
        sheet["F4"] = "整件"
        sheet["F4"].font = Font(italic=True)
        sheet["H6"] = "M001"
        sheet["I6"] = "测试物料"
        sheet["AM6"] = 3
        book.save(source)

        generate_difference(source, target)
        result = openpyxl.load_workbook(target)["机柜差异表"]
        assert result.column_dimensions.get("A") is None
        assert result.sheet_format.defaultColWidth == 1.6


if __name__ == '__main__':
    test_sample()
    test_xlsx_keeps_original_layout_path()
    print('差异表检查通过')
