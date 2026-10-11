"""汇总机柜设计中所有 -MX 工作表的整件物料。"""

from collections import OrderedDict
from copy import copy
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


def output_path_for(source, output_dir=None):
    source = Path(source)
    directory = Path(output_dir) if output_dir else Path(__file__).resolve().parent
    stem = source.stem.replace("机柜设计", "差异表", 1)
    if stem == source.stem:
        stem += "-差异表"
    return directory / f"{stem}.xlsx"


def _headers(ws):
    wanted = {"代号", "名称", "数量", "备注", "更改"}
    found = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8)):
        for cell in row:
            key = "".join(str(cell.value or "").split())
            if key in wanted:
                found[key] = (cell.row, cell.column)
    if set(found) != wanted or len({r for r, _ in found.values()}) != 1:
        raise ValueError(f"{ws.title} 缺少标准表头：{', '.join(sorted(wanted - set(found)))}")
    return found


def _marker_row(ws):
    for row in ws.iter_rows():
        for cell in row:
            if "".join(str(cell.value or "").split()) == "整件" and cell.font.italic:
                return cell.row
    raise ValueError(f"{ws.title} 找不到斜体的“整件”行")


def _copy_rows(source, target, first, last, destination, first_column=1, last_column=None, column_offset=0):
    offset = destination - first
    last_column = last_column or source.max_column
    for row in source.iter_rows(min_row=first, max_row=last,
                                min_col=first_column, max_col=last_column):
        for old in row:
            new = target.cell(old.row + offset, old.column + column_offset)
            new.value = old.value
            _copy_style(old, new)
        old_dim = source.row_dimensions[row[0].row]
        new_dim = target.row_dimensions[row[0].row + offset]
        new_dim.height = old_dim.height
        new_dim.hidden = old_dim.hidden
    for merged in source.merged_cells.ranges:
        if (first <= merged.min_row and merged.max_row <= last
                and first_column <= merged.min_col and merged.max_col <= last_column):
            target.merge_cells(
                start_row=merged.min_row + offset,
                end_row=merged.max_row + offset,
                start_column=merged.min_col + column_offset,
                end_column=merged.max_col + column_offset,
            )


def _copy_style(old, new):
    if old.has_style:
        new.font = copy(old.font)
        new.fill = copy(old.fill)
        new.border = copy(old.border)
        new.alignment = copy(old.alignment)
        new.number_format = old.number_format
        new.protection = copy(old.protection)


def generate_difference(source_path, output_path=None):
    """生成差异表，返回 (输出路径, 不同代号数)。"""
    source_path = Path(source_path)
    if source_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("差异表需要 .xlsx 或 .xlsm 机柜设计文件")
    output_path = Path(output_path) if output_path else output_path_for(source_path)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("输出路径不能与源文件相同")

    workbook = openpyxl.load_workbook(source_path, data_only=False)
    cached = openpyxl.load_workbook(source_path, data_only=True)
    mx_sheets = [ws for ws in workbook.worksheets if ws.title.upper().endswith("-MX")]
    if not mx_sheets:
        raise ValueError("源文件中没有以 -MX 结尾的工作表")

    materials = OrderedDict()
    for ws in mx_sheets:
        headers = _headers(ws)
        marker = _marker_row(ws)
        code_col = headers["代号"][1]
        name_col = headers["名称"][1]
        qty_col = headers["数量"][1]
        for row in range(marker + 1, ws.max_row + 1):
            code_value = ws.cell(row, code_col).value
            if code_value is None or not str(code_value).strip():
                continue
            code = str(code_value).strip()
            raw_qty = cached[ws.title].cell(row, qty_col).value
            try:
                quantity = Decimal(str(raw_qty))
                if not quantity.is_finite():
                    raise InvalidOperation
            except (InvalidOperation, TypeError):
                raise ValueError(f"{ws.title}!{get_column_letter(qty_col)}{row} 数量无效：{raw_qty!r}") from None
            if code not in materials:
                materials[code] = [ws.cell(row, name_col).value, Decimal(0)]
            materials[code][1] += quantity

    template = mx_sheets[0]
    header = _headers(template)
    section = _marker_row(template)
    shift = -5  # 删除模板未使用的 A:E，F 列移到输出 A 列
    output_columns = 52  # A:AZ，原因只占一个普通单元格

    result = openpyxl.Workbook()
    sheet = result.active
    sheet.title = "机柜差异表"
    sheet.sheet_format = copy(template.sheet_format)
    sheet.sheet_properties = copy(template.sheet_properties)
    sheet.page_margins = copy(template.page_margins)
    sheet.print_options = copy(template.print_options)
    sheet.page_setup = copy(template.page_setup)
    sheet.sheet_view.view = "normal"
    sheet.sheet_view.showGridLines = True

    _copy_rows(template, sheet, 2, 3, 2, 6, 56, shift)
    title_stem = source_path.stem.replace("机柜设计", "差异表", 1)
    if title_stem == source_path.stem:
        title_stem += "-差异表"
    title = re.sub(r"-V\d+(?:\.\d+)?$", "", title_stem, flags=re.IGNORECASE)
    sheet.merge_cells(start_row=1, end_row=1, start_column=1, end_column=output_columns)
    sheet.cell(1, 1).value = title
    sheet.cell(1, 1).font = Font(name="宋体", size=14, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for index, (code, (name, quantity)) in enumerate(materials.items(), 1):
        row = 4 + 2 * (index - 1)
        _copy_rows(template, sheet, section + 2, section + 3, row, 6, 56, shift)
        sheet.cell(row, 1).value = index
        sheet.cell(row, header["代号"][1] + shift).value = code
        sheet.cell(row, header["名称"][1] + shift).value = name
        sheet.cell(row, header["数量"][1] + shift).value = int(quantity) if quantity == int(quantity) else float(quantity)
        sheet.cell(row, header["备注"][1] + shift).value = None
        sheet.cell(row, header["更改"][1] + shift).value = f'=IF(AL{row}="","",AH{row}-AL{row})'
        _copy_style(template.cell(section + 2, header["备注"][1]), sheet.cell(row, output_columns))
        sheet.cell(row, output_columns).value = None
        sheet.merge_cells(start_row=row, end_row=row + 1,
                          start_column=output_columns, end_column=output_columns)

    sheet.cell(2, header["备注"][1] + shift).value = "合同数量"
    sheet.cell(2, header["更改"][1] + shift).value = "数量差异"
    _copy_style(template.cell(2, header["备注"][1]), sheet.cell(2, output_columns))
    sheet.cell(2, output_columns).value = "原因"
    sheet.merge_cells(start_row=2, end_row=3,
                      start_column=output_columns, end_column=output_columns)
    sheet.column_dimensions["AZ"].width = 18
    sheet.freeze_panes = None
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.print_area = None
    sheet.print_title_rows = None

    thin = Side(style="thin", color="000000")
    all_lines = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row,
                               min_col=1, max_col=output_columns):
        for cell in row:
            cell.border = all_lines

    output_path.parent.mkdir(parents=True, exist_ok=True)
    result.save(output_path)
    return output_path, len(materials)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="从机柜设计文件生成机柜差异表")
    parser.add_argument("source", help="机柜设计 .xlsx/.xlsm 文件")
    parser.add_argument("--output", help="输出路径，默认在程序目录")
    arguments = parser.parse_args()
    path, count = generate_difference(arguments.source, arguments.output)
    print(f"已生成 {path}，共 {count} 种代号")
